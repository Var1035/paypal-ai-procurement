import openpyxl

wb = openpyxl.load_workbook(r"c:\Users\varun\OneDrive\Desktop\PAYPAL AI\agentmarket_40_company_product_dataset.xlsx")
print("Sheet names:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    print(f"\n--- Sheet: {sheet_name} (rows: {sheet.max_row}, cols: {sheet.max_column}) ---")
    headers = [cell.value for cell in sheet[1]]
    print("Headers:", headers)
    for i in range(2, min(7, sheet.max_row + 1)):
        row_vals = [cell.value for cell in sheet[i]]
        print(f"Row {i}:", row_vals)
